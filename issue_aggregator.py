import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl import load_workbook
import os

def extract_data(filepath):
    """
    個別ファイルからデータを抽出します。
    """
    try:
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active

        data = {
            '発生日': ws['C4'].value,
            '報告者': ws['E4'].value,
            '機種': ws['C5'].value,
            '製番': ws['E5'].value,
            '発生事象': ws['C6'].value,
            '作業区分': ws['C8'].value,
            '暫定対応': ws['E8'].value,
            '関連No': ws['C9'].value
        }
        wb.close()
        return data
    except Exception as e:
        print(f"Error extracting from {filepath}: {e}")
        return None

def append_data(summary_path, data_list):
    """
    集計用ファイルにデータを追記します。
    """
    try:
        if not os.path.exists(summary_path):
            raise FileNotFoundError(f"集計用ファイルが見つかりません: {summary_path}")

        wb = load_workbook(summary_path)
        ws = wb.active

        # 最終行のNoを確認
        last_row = ws.max_row
        while last_row > 1 and ws.cell(row=last_row, column=1).value is None:
            last_row -= 1

        start_no = 1
        if last_row > 1:
            val = ws.cell(row=last_row, column=1).value
            if isinstance(val, int):
                start_no = val + 1

        current_row = last_row + 1
        for data in data_list:
            if data is None:
                continue

            # 転記先列の対応:
            # A: No
            # C: 発生日
            # E: 号機 (機種)
            # G: 発生事象
            # I: 起案者 (報告者)
            # J: 対策方法 (暫定対応)
            # K: 区分 (作業区分)

            ws.cell(row=current_row, column=1).value = start_no
            ws.cell(row=current_row, column=3).value = data['発生日']
            ws.cell(row=current_row, column=5).value = data['機種']
            ws.cell(row=current_row, column=7).value = data['発生事象']
            ws.cell(row=current_row, column=9).value = data['報告者']
            ws.cell(row=current_row, column=10).value = data['暫定対応']
            ws.cell(row=current_row, column=11).value = data['作業区分']

            start_no += 1
            current_row += 1

        wb.save(summary_path)
        wb.close()
        return True
    except Exception as e:
        print(f"Error appending to {summary_path}: {e}")
        return False

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("課題連絡票集計システム")
        self.root.geometry("500x300")

        self.source_files = []
        self.summary_file = ""

        # UI要素の作成
        tk.Label(root, text="1. 個別ファイルを選択してください (複数可)").pack(pady=10)
        tk.Button(root, text="ファイルを選択", command=self.select_sources).pack()
        self.lbl_sources = tk.Label(root, text="選択されていません", fg="gray")
        self.lbl_sources.pack()

        tk.Label(root, text="2. 集計用ファイルを選択してください").pack(pady=10)
        tk.Button(root, text="ファイルを選択", command=self.select_summary).pack()
        self.lbl_summary = tk.Label(root, text="選択されていません", fg="gray")
        self.lbl_summary.pack()

        tk.Button(root, text="集計を実行する", command=self.run_aggregation, bg="lightblue", height=2).pack(pady=20)

    def select_sources(self):
        files = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx")])
        if files:
            self.source_files = list(files)
            self.lbl_sources.config(text=f"{len(self.source_files)} 個のファイルを選択中", fg="black")

    def select_summary(self):
        file = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if file:
            self.summary_file = file
            self.lbl_summary.config(text=os.path.basename(self.summary_file), fg="black")

    def run_aggregation(self):
        if not self.source_files or not self.summary_file:
            messagebox.showwarning("警告", "ファイルを選択してください。")
            return

        all_data = []
        for f in self.source_files:
            data = extract_data(f)
            if data:
                all_data.append(data)

        if all_data:
            success = append_data(self.summary_file, all_data)
            if success:
                messagebox.showinfo("完了", f"{len(all_data)} 件のデータを転記しました。")
            else:
                messagebox.showerror("エラー", "転記中にエラーが発生しました。")
        else:
            messagebox.showwarning("警告", "有効なデータが見つかりませんでした。")

if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
