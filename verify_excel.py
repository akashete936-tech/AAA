import openpyxl

def verify():
    filename = "test_summary.xlsx"
    wb = openpyxl.load_workbook(filename)
    ws = wb.active

    # Check row 2 (first individual data)
    # A: No, C: 発生日, E: 号機, G: 発生事象, I: 起案者, J: 対策方法, K: 区分
    assert ws['A2'].value == 1
    assert ws['C2'].value == '2026/4/1'
    assert ws['E2'].value == 'LVXXX_1号機'
    assert ws['G2'].value == '部品に傷を確認した。'
    assert ws['I2'].value == 'AA氏'
    assert ws['J2'].value == '未'
    assert ws['K2'].value == '定常作業'

    # Check row 3 (second individual data)
    assert ws['A3'].value == 2
    assert ws['C3'].value == '2026/4/2'
    assert ws['E3'].value == 'MVYYY_2号機'
    assert ws['G3'].value == '異音が発生した。'
    assert ws['I3'].value == 'BB氏'
    assert ws['J3'].value == '部品交換'
    assert ws['K3'].value == '非定常作業'

    print("Verification successful!")
    wb.close()

if __name__ == "__main__":
    verify()
