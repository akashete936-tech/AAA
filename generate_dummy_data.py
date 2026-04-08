import openpyxl
from openpyxl import Workbook

def create_individual(filename, date, reporter, model, serial, issue, category, response, related_no):
    wb = Workbook()
    ws = wb.active
    ws['C4'] = date
    ws['E4'] = reporter
    ws['C5'] = model
    ws['E5'] = serial
    ws['C6'] = issue
    ws['C8'] = category
    ws['E8'] = response
    ws['C9'] = related_no
    wb.save(filename)

def create_summary(filename):
    wb = Workbook()
    ws = wb.active
    # Header
    ws['A1'] = 'No'
    ws['B1'] = '入信日'
    ws['C1'] = '発生日'
    ws['D1'] = 'セグメント'
    ws['E1'] = '号機'
    ws['F1'] = '要旨'
    ws['G1'] = '発生事象'
    ws['H1'] = '画像'
    ws['I1'] = '起案者'
    ws['J1'] = '対策方法'
    ws['K1'] = '区分'
    wb.save(filename)

if __name__ == "__main__":
    create_individual("test_individual_1.xlsx", "2026/4/1", "AA氏", "LVXXX_1号機", "H110100001", "部品に傷を確認した。", "定常作業", "未", "None")
    create_individual("test_individual_2.xlsx", "2026/4/2", "BB氏", "MVYYY_2号機", "H220200002", "異音が発生した。", "非定常作業", "部品交換", "123")
    create_summary("test_summary.xlsx")
    print("Dummy files created.")
