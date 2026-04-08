import sys
import argparse
# newspaper3kのlxml依存関係の互換性問題を解決するためのパッチ
try:
    from lxml.html import clean
except ImportError:
    try:
        import lxml_html_clean as clean
        sys.modules['lxml.html.clean'] = clean
    except ImportError:
        pass

from newspaper import Article
from pptx import Presentation
from pptx.util import Inches, Pt

def fetch_article(url):
    """
    URLから記事をダウンロードしてパースし、タイトルと本文を返します。
    """
    try:
        article = Article(url)
        article.download()
        article.parse()
        return {
            'title': article.title,
            'text': article.text,
            'url': url
        }
    except Exception as e:
        print(f"Error fetching {url}: {e}")
        return None

def create_presentation(articles, output_file="news_summary.pptx"):
    """
    記事のリストからパワーポイントファイルを作成します。
    """
    prs = Presentation()

    # タイトルスライド
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "ニュースまとめ"
    subtitle.text = "自動生成プレゼンテーション"

    for art in articles:
        if art is None:
            continue

        # 記事ごとのスライド
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)

        # タイトル
        title_shape = slide.shapes.title
        title_shape.text = art['title']

        # 本文
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True

        # 既存の段落（1つ目の段落）を再利用して、空の段落ができるのを防ぐ
        p = tf.paragraphs[0]

        # 本文の冒頭500文字程度を表示
        summary_text = art['text'][:500] + ("..." if len(art['text']) > 500 else "")
        p.text = summary_text
        p.font.size = Pt(14)

        # 出典URLを小さく追記
        p_url = tf.add_paragraph()
        p_url.text = f"\n出典: {art['url']}"
        p_url.font.size = Pt(10)

    prs.save(output_file)
    print(f"Presentation saved to {output_file}")

def main():
    parser = argparse.ArgumentParser(description="指定したニュースURLからパワーポイントを作成します。")
    parser.add_argument("urls", nargs="+", help="まとめたいニュース記事のURL（複数可）")
    parser.add_argument("-o", "--output", default="news_summary.pptx", help="出力ファイル名 (デフォルト: news_summary.pptx)")

    args = parser.parse_args()

    articles = []
    for url in args.urls:
        print(f"Fetching: {url}")
        art = fetch_article(url)
        if art:
            articles.append(art)

    if articles:
        create_presentation(articles, args.output)
    else:
        print("有効な記事が見つかりませんでした。")

if __name__ == "__main__":
    main()
